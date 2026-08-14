#!/usr/bin/env python3
"""Combine the 14 pertemuan Markdown files into one Word document
(Modul Ajar Database MySQL), with a title page and an auto-updating
Table of Contents field, real Word tables, styled code blocks, and
inline bold/code/italic formatting.

Usage:
    pip install python-docx
    python build_docx.py

Regenerate this after editing any pertemuan-*.md file so the Word
version stays in sync with the Markdown source (source of truth).
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MATERI_DIR = Path(__file__).resolve().parent
OUT_PATH = MATERI_DIR / "Modul_Ajar_Database_MySQL_14_Pertemuan.docx"

FILES = [
    "pertemuan-01-database-rookie.md",
    "pertemuan-02-data-architect.md",
    "pertemuan-03-schema-builder.md",
    "pertemuan-04-crud-ranger.md",
    "pertemuan-05-query-hunter.md",
    "pertemuan-06-data-analyst.md",
    "pertemuan-07-join-master.md",
    "pertemuan-08-query-strategist.md",
    "pertemuan-09-data-guardian.md",
    "pertemuan-10-normalization-master.md",
    "pertemuan-11-database-engineer.md",
    "pertemuan-12-transaction-guardian.md",
    "pertemuan-13-database-wizard.md",
    "pertemuan-14-database-architect.md",
]

BRAND = RGBColor(0x0E, 0x74, 0x90)
DARK = RGBColor(0x1A, 0x1A, 0x1A)


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def add_page_break(doc):
    doc.add_page_break()


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_toc_field(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Klik kanan di sini lalu pilih “Update Field” untuk menampilkan daftar isi."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`.+?`)")


def add_inline_runs(paragraph, text, base_bold=False, base_italic=False, mono=False, size=None):
    """Parse **bold**, *italic*, and `code` inline markers and append runs.
    Recurses into matched spans so nesting (e.g. `code` inside **bold**)
    still renders correctly composed, instead of leaving markers literal."""
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            add_inline_runs(paragraph, part[2:-2], base_bold=True, base_italic=base_italic, mono=mono, size=size)
            continue
        if part.startswith("*") and part.endswith("*") and not part.startswith("**") and len(part) >= 2:
            add_inline_runs(paragraph, part[1:-1], base_bold=base_bold, base_italic=True, mono=mono, size=size)
            continue
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)
        else:
            run = paragraph.add_run(part)
        if base_bold:
            run.bold = True
        if base_italic:
            run.italic = True
        if mono:
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
    return paragraph


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "1E1E2E")
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line.strip() else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_md_table(doc, rows):
    # rows: list of list[str]; rows[1] is the --- separator row (skip it)
    header = rows[0]
    body = [r for i, r in enumerate(rows) if i != 1]
    body = body[1:]  # drop header from body list (already used)
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(header):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        add_inline_runs(p, htext.strip(), base_bold=True)
        set_cell_shading(hdr_cells[i], "0E7490")
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10)
    for row in body:
        cells = table.add_row().cells
        for i, ctext in enumerate(row):
            if i >= n_cols:
                break
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline_runs(p, ctext.strip(), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------------------
# markdown parsing per file
# --------------------------------------------------------------------------
def parse_table_block(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [c.strip() for c in raw.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def render_markdown(doc, text, heading1_page_break=True):
    lines = text.split("\n")
    i = 0
    first_h1 = True
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                add_code_block(doc, code_lines)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("| "):
            rows, new_i = parse_table_block(lines, i)
            if len(rows) >= 2:
                add_md_table(doc, rows)
            i = new_i
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            if level == 1 and not first_h1 and heading1_page_break:
                add_page_break(doc)
            first_h1 = False
            h = doc.add_heading("", level=min(level, 4))
            add_inline_runs(h, heading_text)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline_runs(p, stripped[2:], base_italic=True)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, m.group(2))
            i += 1
            continue

        # plain paragraph — accumulate wrapped lines until blank/special
        para_lines = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4})\s", lines[i].strip()) \
                and not lines[i].strip().startswith(("```", "|", ">", "-", "*")) \
                and not re.match(r"^\d+\.\s", lines[i].strip()) and lines[i].strip() != "---":
            para_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(para_lines))


# --------------------------------------------------------------------------
# document assembly
# --------------------------------------------------------------------------
def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for lvl, size, color in [(1, 20, BRAND), (2, 15, BRAND), (3, 12.5, DARK), (4, 11, DARK)]:
        style = doc.styles[f"Heading {lvl}"]
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.font.name = "Calibri"

    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # ---------------- Title page ----------------
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MODUL AJAR")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = BRAND

    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = title2.add_run("DATABASE MYSQL")
    r2.bold = True
    r2.font.size = Pt(30)
    r2.font.color.rgb = DARK

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = sub.add_run("14 Pertemuan — Kurikulum Outcome-Based Education (OBE)")
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for _ in range(3):
        doc.add_paragraph()

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = desc.add_run(
        "Disusun mengikuti Learning Journey 14 Pertemuan pada Grand Design & Blueprint Gameplay "
        "MYSQL QUEST, memakai studi kasus terpadu ASIA Smart Campus — selaras dengan game "
        "pembelajaran interaktif MYSQL QUEST."
    )
    r4.italic = True
    r4.font.size = Pt(10.5)
    r4.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    for _ in range(6):
        doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run(
        "Catatan: rumusan CPMK/Sub-CPMK dan bobot penilaian pada modul ini mengikuti draft "
        "Grand Design MYSQL QUEST dan perlu diselaraskan dengan RPS resmi program studi."
    )
    rn.font.size = Pt(9)
    rn.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    add_page_break(doc)

    # ---------------- Table of contents ----------------
    doc.add_heading("Daftar Isi", level=1)
    add_toc_field(doc)
    add_page_break(doc)

    # ---------------- Pertemuan content ----------------
    for idx, fname in enumerate(FILES):
        path = MATERI_DIR / fname
        text = path.read_text(encoding="utf-8")
        render_markdown(doc, text, heading1_page_break=(idx > 0))

    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")


if __name__ == "__main__":
    build()
